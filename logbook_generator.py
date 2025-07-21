"""Logbook automation script.

This script parses recipes from ``receitas.docx`` and generates
45 days of shift data following the rotation rules described in
``prompt_codex.md``. The placeholders in ``C4KM_Logbook_modelo.docx``
are filled with ``docxtpl`` and saved as ``logbook_final.docx``.
"""

from datetime import datetime
import random
from collections import defaultdict
from typing import Dict, List

from docxtpl import DocxTemplate
from docx import Document


def get_recipe_from_docx(path: str) -> Dict[str, List[str]]:
    """Extract recipe categories from a DOCX file.

    The document must contain headings with the words ``Protein``/``Proteina``,
    ``Sauce``/``Molho`` and ``Accompaniment``/``Acompanhamento``. Each
    subsequent paragraph is captured as a recipe until a new heading is
    encountered.
    """
    doc = Document(path)
    recipes = {"proteins": [], "sauces": [], "accompaniments": []}
    current = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        lower = text.lower()
        if "prote" in lower and "protein" in lower or "proteína" in lower:
            current = "proteins"
            continue
        if "sauce" in lower or "molho" in lower:
            current = "sauces"
            continue
        if "accompaniment" in lower or "acompanhamento" in lower:
            current = "accompaniments"
            continue
        if current:
            recipes[current].append(text)
    return recipes


def _choose_unique(item_list: List[str], history: Dict[str, int], day: int, gap: int = 8) -> str:
    """Return an item not used in the last ``gap`` days."""
    available = [r for r in item_list if r not in history or day - history[r] >= gap]
    if not available:
        # reset history if we ran out of options
        history.clear()
        available = item_list[:]
    choice = random.choice(available)
    history[choice] = day
    return choice


def get_observation(shift: str) -> str:
    """Return a standard observation based on the shift type."""
    if shift == "Preparation":
        return "General mise en place for service."
    if shift == "Sauce":
        return "Sauce station duties following recipes."
    return "Garnish and finishing touches during service."


def generate_shift_data(recipes: Dict[str, List[str]], days: int = 45) -> List[Dict[str, str]]:
    """Generate shift information for ``days`` days."""
    pattern = ["Preparation", "Sauce", "Garnish"]
    rotation = 5
    histories = {
        "proteins": defaultdict(int),
        "sauces": defaultdict(int),
        "accompaniments": defaultdict(int),
    }
    data = []
    for day in range(1, days + 1):
        shift = pattern[((day - 1) // rotation) % len(pattern)]
        protein = _choose_unique(recipes["proteins"], histories["proteins"], day)
        sauce = _choose_unique(recipes["sauces"], histories["sauces"], day)
        accompaniment = _choose_unique(recipes["accompaniments"], histories["accompaniments"], day)
        data.append({
            "shift": shift,
            "protein": protein,
            "sauce": sauce,
            "accompaniment": accompaniment,
            "obs": get_observation(shift),
        })
    return data


def fill_logbook(template_path: str, output_path: str, days: List[Dict[str, str]]) -> None:
    """Fill the docx template with the provided day data."""
    doc = DocxTemplate(template_path)
    context = {}
    for idx, info in enumerate(days, 1):
        context[f"dia{idx}_shift"] = info["shift"]
        context[f"dia{idx}_protein"] = info["protein"]
        context[f"dia{idx}_sauce"] = info["sauce"]
        context[f"dia{idx}_accompaniment"] = info["accompaniment"]
        context[f"dia{idx}_obs"] = info["obs"]
    doc.render(context)
    doc.save(output_path)


def main() -> None:
    recipes = get_recipe_from_docx("receitas.docx")
    days = generate_shift_data(recipes)
    fill_logbook("C4KM_Logbook_modelo.docx", "logbook_final.docx", days)


if __name__ == "__main__":
    random.seed()
    main()
